"""Export the current table presentation without changing its values or order."""
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def export_table(model, path, hidden_columns):
    columns = [i for i in range(model.columnCount()) if i not in hidden_columns]
    if not columns:
        raise ValueError('没有可导出的显示列')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '主名单'
    edge = Side(style='hair', color='EDF0F5')
    for col, source in enumerate(columns, 1):
        cell = sheet.cell(1, col, model.columns[source][1])
        cell.font = Font(name='Microsoft YaHei UI', size=9, bold=True, color='344054')
        cell.fill = PatternFill('solid', fgColor='EEF2FF')
        cell.alignment = Alignment(horizontal='left', vertical='center')
        width_px = 130 if model.columns[source][0] == 'next_followup_at' else 110
        sheet.column_dimensions[get_column_letter(col)].width = (width_px - 5) / 7
    sheet.row_dimensions[1].height = 28.5  # 38px at 96 DPI
    for row in range(model.rowCount()):
        sheet.row_dimensions[row + 2].height = 15  # 20px at 96 DPI
        for col, source in enumerate(columns, 1):
            value = model.data(model.index(row, source))
            cell = sheet.cell(row + 2, col, value)
            # Keep IDs, lesson lists and ratios exactly as displayed; never formulas.
            cell.data_type = 's'
            cell.number_format = '@'
            cell.font = Font(name='Microsoft YaHei UI', size=7.5, color='344054')
            cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            cell.fill = PatternFill('solid', fgColor='FAFBFF' if row % 2 else 'FFFFFF')
            cell.border = Border(bottom=edge, right=edge)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    target = Path(path)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix='.xlsx', delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        workbook.close()
        if temporary and temporary.exists():
            temporary.unlink()
    return model.rowCount()
