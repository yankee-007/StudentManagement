import json
import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PySide6.QtCore import QCoreApplication
from app.backend import Backend
from tests.profile_fixtures import insert_profile
from app.campaigns import CampaignStore,TEST_TEMPLATE,EXPORT_COLUMNS


class CampaignTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QCoreApplication.instance() or QCoreApplication([])

    def test_cycle_snapshot_feedback_and_export(self):
        with tempfile.TemporaryDirectory() as folder:
            b=Backend(Path(folder)/'test.db')
            flags={f'{p}{i}':'N' for p in ('c','z') for i in range(1,33)}
            source=[]
            with b.db.connect() as conn:
                for i in range(10):
                    sid=str(i)
                    conn.execute('INSERT INTO students(student_id,name,updated_at) VALUES(?,?,?)',(sid,'学员'+sid,'2026-09-24'))
                    insert_profile(conn,sid,'学员'+sid,i,{'微信':'是','合计完课':1,'差的课程':'2','合计作业':0,'差的作业':'1,2'})
                    if i!=9:source.append({'student_id':sid,'flags':dict(flags,c1='T',c2='T' if i==0 else 'F',z1='F',z2='F')})
            b.repo.set_setting('snapshot',json.dumps(source))
            b.repo.update_manual('8',{'status':'请假','exemption_end':'2099-12-31'})
            w=b.workflow
            w.createBatch()
            first=w._batch
            rows=w.store.rows(first)
            self.assertEqual(len(rows),10)
            self.assertEqual(rows[0]['student_id'],'0')
            self.assertEqual(rows[0]['missing_total'],'0/2')
            self.assertEqual(rows[0]['completed_total'],'2/0')
            self.assertEqual(sum(r['eligible'] for r in rows),8)
            w.simulate(True)
            rows=w.store.rows(first)
            self.assertEqual(sum(r['send_state']=='模拟失败' for r in rows),1)
            w.filterRows('pending','')
            sid=w.selected['student_id']
            w.saveDraft(sid,'军训＋考试太忙了')
            resets=[]
            w.tableModel.modelReset.connect(lambda:resets.append(1))
            with patch.object(w,'reload_rows',side_effect=AssertionError('不应刷新全班')):
                self.assertTrue(w.submit('军训＋考试太忙了'))
            self.assertEqual(resets,[])
            next_sid=w.selected['student_id']
            w.saveDraft(next_sid,'尚在编辑')
            count=w.store.mark_unreplied(first)
            # Feedback tracking includes real students regardless of send or exemption status.
            self.assertEqual(count,8)
            self.assertEqual(w.store.mark_unreplied(first),0)
            self.assertEqual(w.store.rows(first,next_sid)[0]['reply_state'],'待反馈')
            self.assertTrue(w.store.rows(first,sid)[0]['feedback'].endswith('军训＋考试太忙了'))
            b.repo.set_setting('snapshot','[]')
            w.createBatch()
            self.assertEqual(len(w.store.rows(w._batch)),10)
            self.assertEqual(w.store.rows(first)[0]['completed_total'],'2/0')
            w.selectBatch(1)
            self.assertFalse(w.canEdit)
            self.assertFalse(w.submit('不能修改历史'))
            reopened=CampaignStore(b.db,b.repo)
            self.assertEqual(reopened.rows(first,next_sid)[0]['draft'],'尚在编辑')
            from app.qt_models import DictTableModel
            from app.xlsx_export import export_table
            from openpyxl import load_workbook
            model=DictTableModel(EXPORT_COLUMNS)
            model.set_rows(reopened.rows(first))
            target=Path(folder)/'export.xlsx'
            export_table(model,target,set())
            book=load_workbook(target)
            self.assertEqual(book.active.max_row,11)
            self.assertEqual([c.value for c in book.active[1]],[x[1] for x in EXPORT_COLUMNS])
            book.close()

    def test_class_isolation_and_binding_guard(self):
        with tempfile.TemporaryDirectory() as folder:
            b=Backend(Path(folder)/'root.db')
            w=b.workflow
            def seed(db,path):
                with db.connect() as conn:
                    conn.execute("INSERT INTO students(student_id,name,updated_at) VALUES('same','新班学生','2026-09-24')")
                    insert_profile(conn,'same','新班学生',1)
                return 1
            with patch('app.workflow.QFileDialog.getOpenFileName',return_value=('dummy.xlsx','')),patch('app.workflow.import_profiles',side_effect=seed):
                w.addClass('隔离测试班')
            self.assertEqual(w.className,'隔离测试班')
            self.assertEqual(b.repo.get('same')['name'],'新班学生')
            w.createBatch()
            self.assertEqual(len(w._batches),1)
            self.assertFalse(b.fetchData())
            w.selectClass(0)
            self.assertIsNone(b.repo.get('same'))
            self.assertEqual(w._batches,[])
            w.selectClass(1)
            self.assertEqual(len(w._batches),1)
            reopened=Backend(Path(folder)/'root.db')
            self.assertEqual(len(reopened.workflow.classes),2)
            reopened.workflow.selectClass(1)
            self.assertEqual(reopened.repo.get('same')['name'],'新班学生')

    def test_new_campaign_fetches_before_snapshot_and_failed_fetch_does_not_create(self):
        with tempfile.TemporaryDirectory() as folder:
            b=Backend(Path(folder)/'test.db')
            with b.db.connect() as conn:
                conn.execute("INSERT INTO students(student_id,name,updated_at) VALUES('001','测试学员','2026-09-24')")
                insert_profile(conn,'001','测试学员',1)
            b.repo.set_setting('snapshot',json.dumps([{'student_id':'001','flags':{'c1':'T','z1':'F'}}]))
            class FakeTask:
                def __init__(self,*args,**kwargs):
                    from unittest.mock import Mock
                    self.succeeded=Mock();self.failed=Mock();self.finished=Mock();self.deleteLater=Mock()
                def start(self):pass
            with patch.object(b.settingsModule,'bindingFor',return_value={'class_id':23,'course_id':2}),patch('app.backend.get_password',return_value='secret'),patch('app.backend.AcquisitionTask',FakeTask):
                b.workflow._classes[0]['term_id']='551'
                b.workflow.registry.set_setting('completion_username','completion')
                b.workflow.registry.set_setting('homework_admin_id','homework')
                b.createCampaign()
                self.assertTrue(b.busy)
                self.assertEqual(b.workflow.store.batches(),[])
                row={'学号':'001','姓名':'测试学员',**{f'{p}{i}':'N' for p in ('C','Z') for i in range(1,33)},'C1':'F','Z1':'T'}
                b._fetch_succeeded(([row],{'双方在读并导出':1}))
            rows=b.workflow.store.rows(b.workflow._batch)
            self.assertEqual(len(b.workflow.store.batches()),1)
            self.assertEqual((rows[0]['courses'],rows[0]['homework']),('1',''))
            self.assertEqual((rows[0]['completed_courses'],rows[0]['completed_homework']),('0','1'))
            self.assertEqual([label for _,label in EXPORT_COLUMNS][2:7],
                             ['未完课次','未完作业','欠交合计','合计完成课程','合计完成作业'])
            self.assertEqual(len(EXPORT_COLUMNS),12)
            self.assertIn(('followup_status','可跟进状态'),EXPORT_COLUMNS)
            with patch.object(b.settingsModule,'bindingFor',return_value={'class_id':23,'course_id':2}),patch('app.backend.get_password',return_value='secret'),patch('app.backend.AcquisitionTask',FakeTask):
                b.createCampaign()
                b._fetch_failed('模拟获取失败')
            self.assertEqual(len(b.workflow.store.batches()),1)

    def test_unmatched_students_do_not_carry_old_learning_into_new_batch(self):
        with tempfile.TemporaryDirectory() as folder:
            b=Backend(Path(folder)/'test.db')
            with b.db.connect() as conn:
                for sid in ('001','002'):
                    conn.execute('INSERT INTO students(student_id,name,updated_at) VALUES(?,?,?)',(sid,'学员'+sid,'2026-09-24'))
                    insert_profile(conn,sid,'学员'+sid,int(sid))
            b.repo.set_setting('snapshot',json.dumps([{'student_id':sid,'flags':{'c1':'F','z1':'F'}} for sid in ('001','002')]))
            with b.db.connect() as conn:
                conn.execute("UPDATE reminder_data SET matched=0,sync_state='本次未获取，保留原值' WHERE student_id='002'")
            b.workflow.createBatch()
            rows={r['student_id']:r for r in b.workflow.store.rows(b.workflow._batch)}
            self.assertEqual(rows['001']['missing_total'],'1/1')
            self.assertEqual(rows['002']['missing_total'],'')
            self.assertEqual(rows['002']['completed_courses'],'')
            self.assertEqual(rows['002']['source_sync'],'')
            self.assertFalse(rows['002']['eligible'])

    def _dashboard_batch(self,folder):
        """在读身份必须走班期名单同步：看板分母只认 roster_status='在读'。"""
        from app.roster_sync import sync_roster
        b=Backend(Path(folder)/'test.db')
        term=dict(termId=564,termNo='P2026175',termName='测试班')
        sync_roster(b.db,term,[dict(student_id=f'P2026175{i:03d}A',name=f'学员{i}',status='在读',
                                    student_type='新生',nickname='',source='接口学员') for i in (1,2)])
        for sid in ('P2026175001A','P2026175002A'):
            b.repo.update_profile_field(sid,'微信','是')
        b.repo.set_setting('snapshot',json.dumps([
            {'student_id':'P2026175001A','flags':{'c1':'T','c2':'T','z1':'T','z2':'F'}},
            {'student_id':'P2026175002A','flags':{'c1':'T','c2':'F','z1':'F','z2':'F'}}]))
        b.workflow.createBatch()
        return b

    def test_dashboard_total_and_cumulative_agree_with_completion_buckets(self):
        with tempfile.TemporaryDirectory() as folder:
            b=self._dashboard_batch(folder)
            data=b.workflow.dashboard
            self.assertEqual(data['version'],3)
            self.assertEqual(data['total'],2)
            self.assertEqual(data['opened'],2)
            self.assertEqual(data['cumulative']['courses'],1)   # 只有 001 完成第 1～2 节
            self.assertEqual(data['cumulative']['courses'],b.workflow.store.dashboard(b.workflow._batch)['courses'][-1]['completed'])
            buckets=data['completion']['courses']
            self.assertEqual([(r['count'],r['people']) for r in buckets],[(2,1),(1,1),(0,0)])
            self.assertEqual(sum(r['people'] for r in buckets),data['total'])
            self.assertEqual([r['cumulative'] for r in buckets],[1,2,2])
            # 无人工是标记时，不把待反馈或已有回复自动计入可跟进。
            self.assertEqual([r.get('followable') for r in buckets],[0,0,0])
            self.assertNotIn('followable',b.workflow.store.dashboard(b.workflow._batch)['completion']['courses'][0])

    def test_old_snapshot_versions_hide_completion_buckets(self):
        with tempfile.TemporaryDirectory() as folder:
            b=self._dashboard_batch(folder)
            batch=b.workflow._batch
            with b.db.connect() as conn:
                stored=json.loads(conn.execute('SELECT data FROM campaign_dashboards WHERE batch_id=?',(batch,)).fetchone()[0])
                stored['version']=2
                conn.execute('UPDATE campaign_dashboards SET data=? WHERE batch_id=?',(json.dumps(stored,ensure_ascii=False),batch))
            b.workflow.refresh_dashboard()
            old=b.workflow.dashboard
            self.assertEqual(old['completion'],{'courses':[],'homework':[]})
            self.assertEqual((old['cumulative']['courses'],old['cumulativeHomework']),('—','—'))
            self.assertIn('旧版快照未保存完课次数分布',old['notice'])
            self.assertNotIn('累计指标无法还原',old['notice'])
            with b.db.connect() as conn:
                stored=json.loads(conn.execute('SELECT data FROM campaign_dashboards WHERE batch_id=?',(batch,)).fetchone()[0])
                stored['version']=1
                conn.execute('UPDATE campaign_dashboards SET data=? WHERE batch_id=?',(json.dumps(stored,ensure_ascii=False),batch))
            b.workflow.refresh_dashboard()
            self.assertIn('累计指标无法还原',b.workflow.dashboard['notice'])
            self.assertEqual(b.workflow.dashboard['courses'][0]['completedRate'],'—')
