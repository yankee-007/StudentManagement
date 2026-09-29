import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from PySide6.QtCore import QCoreApplication
from app.backend import Backend
from app.database import Database
from app.repository import StudentRepository
from app.profile_storage import definitions, set_exemption
from app.profile_fields import BASE_PROFILE_LABELS
from app.importer import import_rows
from tests.test_business_logic import seeded, merged_row
from tests.profile_fixtures import legacy_profiles


class ProfileExtensionTests(unittest.TestCase):
    def test_profile_export_defaults_custom_hidden_and_filter(self):
        from openpyxl import load_workbook
        with seeded(2) as b, tempfile.TemporaryDirectory() as folder:
            p=b.profilesModule
            sid=p.selected['student_id']
            p.autoSaveField(sid,'学习目的','=1+1')
            p.autoSaveField(sid,'画像情况','额外内容')
            for field in p.managedFields:
                p.setFieldVisible(field['field_id'],False)
            p.search(sid)
            keys=[f['key'] for f in p.exportFields if f['selected']]
            path=Path(folder)/'profiles.xlsx'
            self.assertEqual(p.export_to_path(keys,path),1)
            book=load_workbook(path)
            rows=list(book.active.values)
            self.assertEqual(rows[0],('学号','姓名','微信','QQ','电脑','工具安装','所在地区','学习目的'))
            self.assertEqual(rows[1][0],sid)
            self.assertEqual(book.active['H2'].value,'=1+1')
            self.assertEqual(book.active['H2'].data_type,'s')
            book.close()
            p.export_to_path(['name','profile:画像情况'],path)
            book=load_workbook(path)
            self.assertEqual(list(book.active.values)[0],('姓名','画像情况'))
            self.assertEqual(book.active['B2'].value,'额外内容')
            book.close()
            with self.assertRaises(ValueError):p.export_to_path([],path)
            with self.assertRaises(ValueError):p.export_to_path(['unknown'],path)

    def test_column_visibility_and_permanent_extra_deletion(self):
        with seeded(1) as b:
            p=b.profilesModule
            for field in p.managedFields:
                p.setFieldVisible(field['field_id'],False)
            self.assertEqual(p.columnLabels,['学号','姓名'])
            p.refresh()
            self.assertEqual(p.columnLabels,['学号','姓名'])
            p.setFieldVisible('column:profile:微信',True)
            self.assertEqual(p.columnLabels,['学号','姓名','微信'])
            sid=p.selected['student_id']
            self.assertTrue(p.autoSaveField(sid,'画像情况','测试内容'))
            self.assertTrue(p.deleteField('default_2'))
            self.assertNotIn('画像情况',[r['label'] for r in p.fields])
            for db in (Database(b.db.path),Database(b.workflow.registry.db.path)):
                self.assertNotIn('default_2',[r['field_id'] for r in definitions(db)])
                with db.connect() as conn:
                    self.assertEqual(conn.execute("SELECT count(*) FROM profile_field_values WHERE field_id='default_2'").fetchone()[0],0)
            self.assertFalse(p.deleteField('column:student_id'))

    @classmethod
    def setUpClass(cls):cls.app=QCoreApplication.instance() or QCoreApplication([])

    def test_legacy_profile_and_exemption_migrate_once_without_data_loss(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'legacy.db'
            db=Database(path)
            with db.connect() as conn:
                legacy_profiles(conn)
                conn.execute("DELETE FROM settings WHERE key='exemptions_migrated_v1'")
                conn.execute("INSERT INTO students(student_id,name,status,exemption_end,updated_at) VALUES('001','旧姓名','请假','2020-01-02','2020-01-01')")
                fields={'微信':'是','QQ':'否','电脑':'是','工具安装':'待检查','所在地区':'上海','学习目的':'转行','画像情况':'不能丢','开学时间':'九月初','军训时间':'9/2—9/8'}
                conn.execute('INSERT INTO profiles(student_id,name,position,fields) VALUES(?,?,?,?)',('001','旧姓名',1,json.dumps(fields)))
            migrated=Database(path)
            repo=StudentRepository(migrated)
            row=repo.get('001')
            self.assertEqual(row['profile_fields'],fields)
            self.assertEqual(row['exemption_date'],'2020-01-02')
            self.assertTrue(row['exemption_expired'])
            self.assertFalse(row['exemption_active'])
            with migrated.connect() as conn:
                self.assertEqual([r[1] for r in conn.execute('PRAGMA table_info(profiles)')],['student_id','fields'])
                self.assertEqual(set(json.loads(conn.execute('SELECT fields FROM profiles').fetchone()[0])),set(BASE_PROFILE_LABELS))
                self.assertEqual(conn.execute('SELECT count(*) FROM profile_field_values').fetchone()[0],3)
                self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])
                self.assertEqual(conn.execute('PRAGMA foreign_key_list(profiles)').fetchone()['table'],'class_roster')
            set_exemption(migrated,'001','')
            self.assertEqual(StudentRepository(Database(path)).get('001')['exemption_date'],'')

    def test_class_fields_values_isolated_and_date_validation(self):
        with seeded(1) as b:
            p=b.profilesModule
            self.assertTrue(p.addField('方便联系时间','choice','上午\n晚上',False))
            self.assertTrue(p.addField('预计毕业日','date','',True))
            self.assertFalse(p.addField('预计毕业日','text','',True))
            self.assertFalse(p.addField('姓名','text','',True))
            self.assertTrue(p.autoSaveField('P2026169001A','方便联系时间','晚上'))
            self.assertFalse(p.autoSaveField('P2026169001A','方便联系时间','中午'))
            self.assertTrue(p.autoSaveField('P2026169001A','预计毕业日','2028-06-30'))
            self.assertFalse(p.autoSaveField('P2026169001A','预计毕业日','not-a-date'))
            self.assertNotIn('方便联系时间',p.columnLabels)
            field=next(r for r in p.extraFields if r['name']=='方便联系时间')
            p.setFieldVisible(field['field_id'],True)
            self.assertIn('方便联系时间',p.columnLabels)
            b.workflow._classes.append(dict(name='新班',path=str(b.db.path.parent/'second.db')))
            b.workflow.selectClass(1)
            self.assertNotIn('预计毕业日',[r['name'] for r in definitions(b.db)])
            with b.db.connect() as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM profile_field_values').fetchone()[0],0)
            b.workflow.selectClass(0)
            self.assertEqual(b.repo.get('P2026169001A')['profile_fields']['预计毕业日'],'2028-06-30')

    def test_exemption_shared_inclusive_expired_reactivated_and_snapshots(self):
        with seeded(1) as b:
            sid='P2026169001A'
            today=date.today().isoformat()
            tomorrow=(date.today()+timedelta(days=1)).isoformat()
            p=b.profilesModule
            p.selectRow(0)
            self.assertTrue(p.saveEditorField(p.selected['_record_key'],'免催日期',today))
            import_rows(b.db,[merged_row()])
            b.workflow.createBatch()
            first=b.workflow._batch
            self.assertEqual(b.workflow.selected['exemption_date'],today)
            self.assertFalse(b.workflow.selected['current_eligible'])
            with patch.object(b,'chooseDate',return_value=tomorrow):b.workflow.setLeave()
            self.assertEqual(p.selected['exemption_date'],tomorrow)
            b.workflow.createBatch()
            self.assertEqual(b.workflow.store.rows(first)[0]['exemption_date'],today)
            b.workflow.clearLeave()
            self.assertEqual(p.selected['exemption_date'],'')
            b.workflow.simulate(False)
            self.assertEqual(b.workflow.selected['send_state'],'模拟成功')
            # A date expiring yesterday remains stored but no longer suppresses sending.
            with b.db.connect() as conn:
                conn.execute('INSERT INTO exemptions VALUES(?,?)',(sid,(date.today()-timedelta(days=1)).isoformat()))
            b.workflow.createBatch()
            row=b.workflow.selected
            self.assertTrue(row['exemption_expired'])
            self.assertTrue(row['current_eligible'])
            self.assertTrue(row['exemption_text'].startswith('至'))
            model=b.workflow.tableModel
            index=next(i for i,(k,_) in enumerate(model.columns) if k=='exemption_text')
            self.assertTrue(model.data(model.index(0,index),model.ExpiredCellRole))
            b.workflow.simulate(False)
            self.assertEqual(b.workflow.selected['send_state'],'模拟成功')


if __name__=='__main__':unittest.main()
