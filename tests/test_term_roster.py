import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from PySide6.QtWidgets import QApplication
from app.backend import Backend
from app.database import Database
from app.term_roster import arrange_students, TermRosterStore, COLUMNS
from tests.profile_fixtures import insert_profile

TERM = dict(termId=551, termNo='P2026169', termName='编程169期')


def student(n, letter='A', name='测试'):
    return dict(student_id=f'P2026169{n:03d}{letter}', name=name)


class TermRosterTest(unittest.TestCase):
    def test_shared_class_selection_never_shows_another_terms_cache(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            backend = Backend(Path(tmp) / 'main.db')
            with backend.db.connect() as conn:
                conn.execute("INSERT INTO students(student_id,name,updated_at) VALUES('legacy001','示例旧班','2026-10-06')")
                insert_profile(conn, 'legacy001', '示例旧班', 1)
            module = backend.termsModule
            # Refreshing remote terms must leave the unrelated legacy class selected.
            self.assertIsNone(module._accept('terms', [TERM]))
            self.assertEqual(backend.workflow.classIndex, 0)
            self.assertEqual(module.termIndex, -1)
            module.store.save_lessons(TERM['termId'], [dict(resource_id='first', label='01【Python】')], 'first')
            module.store.save(TERM, [student(1)], 'first')
            with patch.object(module, '_start') as start:
                backend.workflow.selectClass(1)
                module.activate()
                self.assertEqual(module.termIndex, 0)
                self.assertEqual(module.visibleCount, 1)
                backend.workflow.selectClass(0)
                module.activate()
                self.assertEqual(module.termIndex, -1)
                self.assertEqual(module.visibleCount, 0)
                self.assertEqual(module.lessons, [])
                self.assertIn('未关联平台班期', module.notice)
                start.assert_not_called()

    def test_sort_gaps_and_letters(self):
        rows = arrange_students(TERM['termNo'], [student(10,'E'), student(2,'D'), student(4,'B')])
        self.assertEqual([r['ordinal'] for r in rows], ['',2,'',4,'','','','','',10])
        self.assertEqual(rows[0]['student_id'], 'P2026169001A')
        self.assertEqual(rows[2]['student_id'], 'P2026169003D')
        self.assertEqual(rows[8]['student_id'], 'P2026169009B')
        self.assertEqual(rows[2]['name'], '')
        self.assertEqual(rows[2]['source'], '缺号补位')
        # Same ordinal, different suffixes: neither real student may be dropped.
        self.assertEqual(len(arrange_students(TERM['termNo'], [student(1,'A'),student(1,'B')])), 2)
        self.assertEqual(arrange_students(TERM['termNo'], []), [])

    def test_validation_and_atomic_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(Path(tmp)/'test.db')
            store = TermRosterStore(db)
            store.save(TERM,[student(1),student(3,'D')],'first')
            for invalid in ([student(1),student(1)], [dict(student_id='P2026175001A',name='')], []):
                with self.assertRaises(ValueError): store.save(TERM,invalid,'bad')
                self.assertEqual(store.load(551)['resource_id'], 'first')
            store.save(TERM,[student(1),student(2,'E'),student(3,'D')],'first')
            self.assertEqual(store.load(551)['rows'][1]['student_id'],'P2026169002E')
            self.assertEqual(store.load(551)['rows'][1]['source'],'接口学员')
            self.assertEqual(TermRosterStore(Database(db.path)).load(551),store.load(551))
            store.save(dict(termId=564,termNo='P2026175'),[], 'other')
            self.assertEqual(len(store.load(551)['rows']),3)
            with db.connect() as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM profiles').fetchone()[0],0)
                self.assertEqual(conn.execute('SELECT count(*) FROM students').fetchone()[0],0)

    def test_module_cache_selection_and_first_lesson(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            backend = Backend(Path(tmp)/'main.db')
            module = backend.termsModule
            module._accept('terms',[TERM])
            module._request_term = TERM
            module._accept('lessons',[
                dict(resource_id='pre',label='VIP预热课'),
                dict(resource_id='first',label='01【Python核心语法】'),
                dict(resource_id='eleventh',label='11【Excel】')])
            self.assertEqual(module.lessonIndex,1)
            module._resource='first'
            module._accept('students',[student(1),student(3,'D')])
            module.filterRows('缺号补位')
            self.assertEqual(module.visibleCount,1)
            self.assertIn('真实学员 2 人',module.summary)
            reloaded=Backend(Path(tmp)/'main.db').termsModule
            self.assertEqual(reloaded.visibleCount,3)
            self.assertEqual(reloaded.terms[0]['termId'],551)
            # Known cached course wins, ambiguous titles never choose preheating.
            module._accept('lessons',[dict(resource_id='pre',label='VIP预热课')])
            self.assertEqual(module.lessonIndex,-1)

    def test_api_payload_html_and_partial_rejection(self):
        from app.acquisition.completion import CompletionClient
        client = CompletionClient.__new__(CompletionClient)
        client.request = Mock(return_value='<select id="resourceId"><option value="">所有</option><option value="a">01【课程】</option></select>')
        self.assertEqual(client.lessons('551'),
                         [dict(resource_id='a',label='01【课程】')])
        client.request.return_value='<html>login</html>'
        with self.assertRaises(ValueError): client.lessons('551')
        client._single_page = Mock(return_value=[dict(studentNo='P2026169001A',realname='测试')])
        self.assertEqual(client.students('551','a'),[dict(student(1),status='',student_type='',nickname='')])
        form=client._single_page.call_args.kwargs['form']
        self.assertEqual((form['pageSize'],form['pageNum'],form['resourceId']),(500,1,'a'))
        self.assertEqual(form['status'],'')
        client._single_page.return_value=[dict(studentNo='P2026169001A',realname='',nickname='昵称甲',
                                         xeNickname='关联用户名',status=0,studentType=2)]
        self.assertEqual(client.students('551','a'),[dict(student(1,name=''),status='在读',student_type='冻转',nickname='昵称甲')])
        client._single_page.return_value=[dict(studentNo='P2026169001A',realname='测试',status='2',studentType='1')]
        self.assertEqual(client.students('551','a')[0]['status'],'已退课')
        self.assertEqual(client.students('551','a')[0]['student_type'],'重修')
        client._single_page.return_value=[dict(studentNo='P2026169001A',unexpected='名字')]
        with self.assertRaises(ValueError):client.students('551','a')
        client._single_page.return_value=[dict(studentNo='',realname='测试')]
        with self.assertRaises(ValueError):client.students('551','a')
        del client.request
        del client._single_page
        response=Mock(status_code=200,url='https://xs.jihuaxueyuan.com/list')
        response.json.return_value={'code':0,'rows':[{}],'total':2}
        client.session=Mock()
        client.session.request.return_value=response
        with self.assertRaises(ValueError):client._single_page('POST','/list',form={})
        response.json.return_value={'code':500,'rows':[]}
        with patch('app.acquisition.completion.time.sleep'):
            with self.assertRaises(ValueError):client.request('POST','/list',form={})
        response.json.return_value={'code':403,'msg':'未登录'}
        client.refreshed=True
        with self.assertRaises(ValueError):client.request('POST','/list',form={})

    def test_cache_first_and_forced_refresh_chain(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            module = Backend(Path(tmp)/'main.db').termsModule
            lesson = [dict(resource_id='first',label='01【Python】')]
            with patch.object(module, '_start') as start:
                module.activate()
                start.assert_called_once_with('terms')
                self.assertEqual(module._accept('terms',[TERM]),'lessons')
                module._request_term=TERM
                self.assertEqual(module._accept('lessons',lesson),'students')
                module._resource='first'
                module._accept('students',[student(1)])
                start.reset_mock()
                module.selectTerm(0)
                module.activate()
                start.assert_not_called()
                restarted=Backend(Path(tmp)/'main.db').termsModule
                with patch.object(restarted,'_start') as restarted_start:
                    restarted.activate()
                    restarted_start.assert_not_called()
                    self.assertEqual(restarted.lessons,lesson)
                    self.assertEqual(restarted.lessonIndex,0)
                module.refreshAll()
                start.assert_called_once_with('terms')
                self.assertEqual(module._accept('terms',[TERM]),'lessons')
                self.assertEqual(module._accept('lessons',lesson),'students')
                module._accept('students',[student(1),student(2)])
                self.assertEqual(module.visibleCount,2)
                self.assertFalse(module._force_roster)
                start.reset_mock()
                module.activate()
                start.assert_not_called()

    def test_six_columns_blank_placeholders_legacy_cache_and_export(self):
        import json
        from openpyxl import load_workbook
        from app.qt_models import DictTableModel
        from app.xlsx_export import export_table
        app = QApplication.instance() or QApplication([])
        self.assertEqual([label for _,label in COLUMNS], ['序号','学员学号','学员姓名','状态','类型','昵称'])
        with tempfile.TemporaryDirectory() as tmp:
            db=Database(Path(tmp)/'test.db')
            store=TermRosterStore(db)
            rows=store.save(TERM,[dict(student(1),status='在读',student_type='新生',nickname='测试昵称'),student(3)],'first')['rows']
            self.assertEqual(rows[0]['nickname'],'测试昵称')
            for key,_ in COLUMNS:
                self.assertEqual(rows[1][key], 'P2026169002A' if key=='student_id' else '已退课' if key=='status' else '')
            # Old stored placeholders also become blank, without network refresh.
            legacy=[dict(student_id='P2026169002A',name='',ordinal=2,source='缺号补位')]
            with db.connect() as conn:
                conn.execute('UPDATE term_rosters SET rows_json=? WHERE term_id=?',(json.dumps(legacy),'551'))
            self.assertEqual(store.load(551)['rows'][0]['ordinal'],'')
            model=DictTableModel(COLUMNS)
            model.set_rows(rows)
            path=Path(tmp)/'roster.xlsx'
            export_table(model,path,[])
            book=load_workbook(path)
            try:
                sheet=book.active
                self.assertEqual([c.value for c in sheet[1]],[label for _,label in COLUMNS])
                self.assertEqual([c.value for c in sheet[3]],[None,'P2026169002A',None,'已退课',None,None])
            finally:
                book.close()

    def test_partial_empty_cache_and_ambiguous_course(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            module=Backend(Path(tmp)/'main.db').termsModule
            module._accept('terms',[TERM,dict(termId=564,termNo='P2026175',termName='175期')])
            module._request_term=TERM
            self.assertIsNone(module._accept('lessons',[dict(resource_id='ambiguous',label='Python课程')]))
            self.assertIn('未能唯一识别',module.notice)
            with patch.object(module,'_start') as start:
                module.selectLesson(0)
                start.assert_not_called()
            module._resource='ambiguous'
            module._accept('students',[])
            # An explicitly successful empty roster is cached, not fetched forever.
            with patch.object(module,'_start') as start:
                module.activate()
                start.assert_not_called()
                module.selectTerm(1)
                start.assert_called_once_with('lessons')
                self.assertEqual(module.visibleCount,0)
            # Old installation: roster exists but lessons were never persisted.
            with module.registry.db.connect() as conn:
                conn.execute('DELETE FROM term_lessons WHERE term_id=?',('551',))
            with patch.object(module,'_start') as start:
                module.selectTerm(0)
                start.assert_called_once_with('lessons')
                module._request_term=TERM
                self.assertIsNone(module._accept('lessons',[dict(resource_id='ambiguous',label='01【Python】')]))
    def test_course_selection_never_changes_roster_source(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            module=Backend(Path(tmp)/'main.db').termsModule
            module._accept('terms',[TERM])
            module._request_term=TERM
            module._accept('lessons',[dict(resource_id='warmup',label='预热课'),
                dict(resource_id='first',label='01【Python】'),dict(resource_id='second',label='02【Python】')])
            with patch.object(module,'_start') as start:
                module.selectLesson(2)
                start.assert_not_called()
            self.assertEqual(module.lessonIndex,2)
            class FakeTask:
                def __init__(self,*args,**kwargs):
                    from unittest.mock import Mock
                    self.action=args[0];self.resource_id=kwargs['resource_id']
                    self.succeeded=Mock();self.failed=Mock();self.finished=Mock();self.deleteLater=Mock()
                def start(self):pass
            module.registry.set_setting('completion_username','test')
            with patch('app.term_module.get_password',return_value='secret'),patch('app.term_module.AcquisitionTask',FakeTask):
                module.fetchStudents()
                self.assertEqual(module._task.resource_id,'first')
                module._cleanup()
            module._load_cache()
            self.assertEqual(module.lessonIndex,2)
            self.assertEqual(module._choose_lesson(module.lessons),1)

if __name__ == '__main__': unittest.main()
