import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from app.backend import STUDENT_COLUMNS
from app.qt_models import DictTableModel
from app.xlsx_export import export_table


class ExportTests(unittest.TestCase):
    def test_display_values_and_hidden_columns(self):
        model = DictTableModel(STUDENT_COLUMNS)
        model.set_rows([{'student_id': '001', 'name': '=1+1',
                        'pending_courses_text': '第1节完课、第2节完课',
                        'pending_homework_text': '第3节作业',
                        'pending_courses': ['1', '2'], 'pending_homework': ['3'],
                        'status': '请假', 'next_followup_at': '2026-10-01'}])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'list.xlsx'
            export_table(model, path, {0})
            book = load_workbook(path)
            sheet = book.active
            self.assertEqual(list(sheet.values)[1], ('=1+1', '1,2', '3', '2/1', '请假', '2026-10-01'))
            self.assertEqual(sheet['A2'].data_type, 's')
            self.assertEqual(sheet['A2'].font.sz, 7.5)
            self.assertEqual(sheet.row_dimensions[2].height, 15)
            book.close()
            export_table(model, path, set())
            book = load_workbook(path)
            self.assertEqual(book.active['A2'].value, '001')
            book.close()
